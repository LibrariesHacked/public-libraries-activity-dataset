"""Audit original returns without treating missing cells as reported zeroes."""

import argparse
import csv
from collections import Counter
from datetime import timedelta
from itertools import combinations
import math
from pathlib import Path
from statistics import median

import openpyxl

import rotate_activity_data as rotation


OUTPUT = Path('data/activity_audit.csv')
FIELDS = ['Period', 'Authority code', 'Authority name', 'Dataset', 'Field', 'Check',
          'Severity', 'Known issue', 'Reported', 'Comparator', 'Source cells', 'Notes']
MAPPINGS = Path('data/activity_column_mappings.csv')
RETURNS = Path('data/activity_audit_coverage.csv')
EXPLANATIONS = Path('data/activity_respondent_explanations.csv')
DATASETS = {
    'digital_loans': 'loans', 'total_active_members': 'users',
    'active_members_11_under': 'users', 'active_members_12_17': 'users',
    'active_members_adults': 'users', 'hours_public_computers': 'computer_usage',
    'wifi_sessions': 'wifi_sessions', 'click_and_collect': 'click_and_collect',
    'mobile_libraries': 'visits', 'home_delivery': 'visits',
}
BREAKDOWNS = {
    'total_active_members': ('active_members_11_under', 'active_members_12_17', 'active_members_adults'),
    'total_physical_events': ('physical_events_11_under', 'physical_events_12_17', 'physical_events_adults', 'physical_events_all_ages'),
    'total_digital_events': ('digital_events_11_under', 'digital_events_12_17', 'digital_events_adults', 'digital_events_all_ages'),
    'total_attendees_physical_events': ('physical_attendees_11_under', 'physical_attendees_12_17', 'physical_attendees_adults'),
    'total_attendees_digital_events': ('digital_attendees_11_under', 'digital_attendees_12_17', 'digital_attendees_adults'),
    'total_physical_book_issues': ('loans_11_under', 'loans_12_17', 'loans_adult'),
    'total_physical_audiobook_issues': ('loans_11_under_digital', 'loans_12_17_digital', 'loans_adult_digital'),
    'total_ebook_issues': ('ebooks_11_under', 'ebooks_12_17', 'ebooks_adult'),
    'total_digital_audiobook_issues': ('digital_audiobook_issues_11_under', 'digital_audiobook_issues_12_17', 'digital_audiobook_issues_adult'),
}
AUXILIARY = {
    'events_physical_audience_summary': ('events', 'total_physical_events'),
    'events_digital_audience_summary': ('events', 'total_digital_events'),
    'attendance_physical_age_summary': ('attendance', 'total_attendees_physical_events'),
    'attendance_digital_age_summary': ('attendance', 'total_attendees_digital_events'),
}


def explanation_columns(config):
    if config['mapper'] == 'published_labels':
        return ['Ebook & Eaudio Data Collection',
                'Please Tell Us Why You Have No Data To Enter Here Eg You Do Not Have A Mobile Library Service',
                'Please Tell Us Why You Have No Data To Enter Here Eg You Do Not Collate Data For Wifi']
    return ['Q22'] if config['member_prefix'] == 'Q7' else ['Q13']


def numeric(value):
    if value in (None, '') or isinstance(value, bool):
        return None
    if isinstance(value, timedelta):
        value = value.total_seconds() / 86400
    try:
        number = float(value)
        return number if math.isfinite(number) else None
    except (TypeError, ValueError):
        return None


def column_mappings(year, config):
    columns = []
    if config['mapper'] == 'question_codes':
        for field, group in config['member_groups'].items():
            columns.append(('users', field, 'Annual', f"{config['member_prefix']}_{group}_1"))
        for dataset, groups in config['question_code_groups'].items():
            for group, field in enumerate(groups, 1):
                prefix = config['metric_prefixes'][dataset]
                dataset_name = DATASETS.get(field, DATASETS.get(dataset, dataset))
                for month, code in zip(rotation.MONTHS, config['month_codes'][dataset]):
                    columns.append((dataset_name, field, month, f'{prefix}_{group}_{code}'))
                columns.append((dataset_name, field, 'Annual', f'{prefix}_{group}_Total'))
        for field, metadata in AUXILIARY.items():
            dataset = metadata[0]
            prefix = config['metric_prefixes'][dataset]
            delivery = 'Physical' if 'physical' in field else 'Digital'
            suffix = 'Audience_Total' if dataset == 'events' else 'All_Totals'
            columns.append((dataset, field, 'Annual', f'{prefix}_{delivery}_{suffix}'))
        if year == '2025/2026':
            columns.append(('users', 'active_members_vision_print_impairment', 'Annual', 'Q18a'))
    else:
        for field, header in rotation.MEMBER_LABELS.items():
            columns.append(('users', field, 'Annual', header))
        for field, label in rotation.MONTHLY_LABELS.items():
            dataset = next((name for total, detail in BREAKDOWNS.items()
                            if field in (total, *detail)
                            for name in ['users' if total == 'total_active_members' else
                                         'attendance' if 'attendees' in total else
                                         'events' if 'events' in total else 'loans']), 'visits')
            dataset = DATASETS.get(field, dataset)
            for index, month in enumerate(rotation.MONTHS):
                calendar_year = rotation.financial_year_start(year) + (index + 3) // 12
                header = f'{label} - {rotation.LEGACY_MONTH_NAMES[index]} {calendar_year}'
                columns.append((dataset, field, month, header))
                if field in ('loans_adult', 'loans_11_under', 'loans_12_17'):
                    columns.append(('loans', f'{field}_digital', month, f'{header}_digital'))
    for field, header in config['inventory_fields'].items():
        columns.append(('computer_inventory', field, 'Annual', header))
    return columns


def annual(series):
    if series['annual'] is not None:
        return series['annual']
    values = [value for value in series['monthly'] if value is not None]
    return sum(values) if values else None


def comparable(series):
    supplied = [index for index, value in enumerate(series['monthly']) if value is not None and value > 0]
    return (series['annual'] is not None or len(supplied) >= 10 or
            (len(supplied) >= 3 and all(index in (2, 5, 8, 11) for index in supplied)))


def respondent_context(return_row, field, check, severity, notes):
    explanation = return_row.get('outreach_explanation', '')
    if explanation and field in ('click_and_collect', 'mobile_libraries', 'home_delivery'):
        notes += f' Respondent explanation: {explanation}'
        if (return_row['name'] == 'Rutland' and return_row['year'] == '2024/2025'
                and field == 'home_delivery' and check in ('monthly_outlier', 'year_change')):
            severity = 'information'
            notes += ' The stated service change explains the fall; not promoted as a suspected typo.'
    return severity, notes


def annual_formula_evidence(formula, expected, numeric_text_sum=0, difference=None):
    if not isinstance(formula, str) or not formula.startswith('='):
        return 'Annual total is directly entered, not a formula summing these months.'
    if formula in (expected, f'=ROUND({expected[1:]},0)'):
        if numeric_text_sum > 0 and difference is not None and abs(difference - numeric_text_sum) < 1:
            return (f'Annual SUM ignores numeric text month cells totalling {numeric_text_sum:,.0f}, '
                    'which explains the entire difference. Numeric month values are preserved by the importer; '
                    'the published annual value remains unchanged and flagged.')
        return ('Annual formula sums the intended month cells, but its cached result disagrees with their numeric sum. '
                'Possible stale cached result; no automatic correction inferred.')
    return f'Annual formula {formula} differs from the expected month sum {expected}; inspect its references.'


def audit():
    findings, mappings, coverage, returns, explanations = [], [], [], [], []
    population = rotation.load_population_lookup(rotation.POPULATION)
    services = rotation.load_library_services(rotation.LIBRARY_SERVICES)
    authorities = rotation.load_authorities_lookup(
        rotation.AUTHORITIES, services, population,
        rotation.load_nearest_neighbours_lookup(rotation.NEAREST_NEIGHBOURS))

    def add(return_row, field, check, reported, comparator, notes, severity='suspicious', cells=None):
        series = return_row['series'][field]
        known_rule = rotation.DATA_QUALITY_ANOMALIES.get(
            (return_row['year'], series['dataset'], return_row['code']), {})
        known = bool(known_rule.get('all') or known_rule.get('total') or
                     any(match in field for match in known_rule.get('fields', {})) or
                     any(str(value) in known_rule.get('values', {}) or
                         str(int(value)) in known_rule.get('values', {})
                        for value in [series['annual'], *series['monthly']]
                         if value is not None))
        severity, notes = respondent_context(return_row, field, check, severity, notes)
        findings.append(dict(zip(FIELDS, [
            return_row['year'], return_row['code'], return_row['name'], series['dataset'],
            field, check, severity, 'yes' if known else '', reported, comparator,
            f"{return_row['sheet']} row {return_row['row']}: " + ', '.join(cells or series['headers']), notes,
        ])))

    for year, config in rotation.YEAR_SOURCES.items():
        workbook = openpyxl.load_workbook(config['workbook'], read_only=True, data_only=True)
        worksheet = workbook[config['worksheet']]
        formula_workbook = openpyxl.load_workbook(config['workbook'], read_only=True, data_only=False)
        formula_worksheet = formula_workbook[config['worksheet']]
        first_rows = list(worksheet.iter_rows(max_row=2, values_only=True))
        headers = first_rows[0]
        columns = column_mappings(year, config)
        if len(set(headers)) != len(headers):
            raise ValueError(f'{year}: duplicate source headers')
        for dataset, field, month, header in columns:
            if header not in headers:
                raise ValueError(f'{year} {field}: missing header {header}')
            description = str(first_rows[1][headers.index(header)]) if config['mapper'] == 'question_codes' else header
            if month != 'Annual' and config['mapper'] == 'question_codes':
                label = description.rsplit(' - ', 1)[-1].lower()
                if not label.startswith(month[:3]):
                    raise ValueError(f'{year} {header}: expected {month}, got {label}')
            mappings.append({'Period': year, 'Dataset': dataset, 'Field': field,
                             'Month': month, 'Source column': header,
                             'Source description': description,
                             'Mapping note': 'Annual label uses the 2025/26 demographic order; use the verified monthly question group and formula, not this label.'
                             if year == '2024/2025' and header.endswith('_Total') and
                             any(token in field for token in ('adult', '11_under', '12_17')) else ''})
        seen = set()
        source_rows = zip(worksheet.iter_rows(min_row=2, values_only=True),
                  formula_worksheet.iter_rows(min_row=2, values_only=True))
        for row_number, (values, formulas) in enumerate(source_rows, 2):
            source = dict(zip(headers, values))
            authority = source.get('Q1' if config['mapper'] == 'question_codes' else 'Library Authority (Upper Tier Local Authority)')
            if not authority or str(authority).lower().startswith('library service') or str(authority).startswith('{'):
                continue
            name = rotation.AUTHORITY_ALIASES.get(authority, authority)
            if name not in authorities:
                raise ValueError(f'{year}: unmapped authority {name}')
            code = authorities[name]['gss-code']
            if code in seen:
                raise ValueError(f'{year}: duplicate authority {name}')
            seen.add(code)
            return_row = {'year': year, 'name': name, 'code': code, 'sheet': worksheet.title,
                          'row': row_number, 'series': {}}
            for header in explanation_columns(config):
                text = source.get(header)
                if text not in (None, '', 0, '0', 'nan'):
                    explanation = ' '.join(str(text).split())
                    if header in ('Q22', 'Q13'):
                        return_row['outreach_explanation'] = explanation
                    explanations.append({'Period': year, 'Authority code': code,
                                         'Authority name': name, 'Source column': header,
                                         'Source row': row_number, 'Explanation': explanation})
            nonempty, numeric_count, zero_count = 0, 0, 0
            for dataset, field, month, header in columns:
                series = return_row['series'].setdefault(field, {
                    'dataset': dataset, 'annual': None, 'monthly': [None] * 12, 'headers': [],
                })
                series['headers'].append(header)
                raw = source.get(header)
                value = numeric(raw)
                nonempty += raw not in (None, '')
                numeric_count += value is not None
                zero_count += value == 0
                if month == 'Annual':
                    series['annual'] = value
                    series['annual_formula'] = formulas[headers.index(header)]
                else:
                    series['monthly'][rotation.MONTHS.index(month)] = value
                    series.setdefault('monthly_columns', []).append(headers.index(header) + 1)
                    if isinstance(raw, str) and value is not None:
                        series['numeric_text_sum'] = series.get('numeric_text_sum', 0) + value
                if value is not None and (value < 0 or not value.is_integer()):
                    add(return_row, field, 'invalid_count', value, '',
                        f'{header} contains {value}; counts should be non-negative whole numbers.', cells=[header])
                elif raw not in (None, '') and value is None:
                    add(return_row, field, 'non_numeric', str(raw), '',
                        'Non-numeric response in a numeric column; not silently treated as zero.',
                        severity='information', cells=[header])
                if isinstance(raw, timedelta):
                    add(return_row, field, 'excel_duration', value, str(raw),
                        'Excel duration formatting: underlying day count interpreted as the submitted count, '
                        'matching the importer. Confirm this is formatting, not a time-unit return.',
                        severity='information', cells=[header])
            returns.append(return_row)
            for series in return_row['series'].values():
                monthly_columns = series.get('monthly_columns', [])
                if monthly_columns:
                    first_column = openpyxl.utils.get_column_letter(min(monthly_columns))
                    last_column = openpyxl.utils.get_column_letter(max(monthly_columns))
                    series['expected_formula'] = f'=SUM({first_column}{row_number}:{last_column}{row_number})'
            coverage.append({'Period': year, 'Authority code': code, 'Authority name': name,
                             'Mapped cells': len(columns), 'Nonempty cells': nonempty,
                             'Numeric cells': numeric_count, 'Zero cells': zero_count,
                             'Missing cells': len(columns) - nonempty,
                             'Respondent notes': ' | '.join(
                                 f'{header}: {value}' for header, value in source.items()
                                 if header in explanation_columns(config)
                                 and value not in (None, '')
                                 and header not in ('Q1', 'Library Details', 'Library Authority (Upper Tier Local Authority)'))})
        workbook.close()
        formula_workbook.close()

    for return_row in returns:
        series_by_field = return_row['series']
        for field, series in series_by_field.items():
            monthly, total = series['monthly'], series['annual']
            positive = [value for value in monthly if value is not None and value > 0]
            if total is not None and positive and abs(total - sum(value or 0 for value in monthly)) > 1:
                monthly_sum = sum(value or 0 for value in monthly)
                add(return_row, field, 'annual_month_sum', total, monthly_sum,
                    f'Published annual total {total:,.0f}; supplied month sum {monthly_sum:,.0f}; '
                    f'difference {total - monthly_sum:+,.0f}. No replacement inferred. '
                    + annual_formula_evidence(series.get('annual_formula'), series['expected_formula'],
                                              series.get('numeric_text_sum', 0), monthly_sum - total))
            if len(positive) >= 6:
                typical = median(positive)
                for index, value in enumerate(monthly):
                    if value is not None and value > 0 and typical >= 100 and (value >= typical * 5 or value <= typical * 0.15):
                        add(return_row, field, 'monthly_outlier', value, typical,
                            f'{rotation.MONTHS[index].title()} {value:,.0f} versus median nonzero month {typical:,.0f} '
                            f'({value / typical:.2f} times). Could be a typo, changed scope or genuine activity; confirm with service.')
                zero_months = [rotation.MONTHS[index] for index, value in enumerate(monthly) if value == 0]
                if zero_months and field.startswith('total_'):
                    add(return_row, field, 'zero_months', 0, typical,
                        f'Explicit zero in {", ".join(zero_months)} despite {len(positive)} positive months; '
                        'check whether zero means no activity or missing data.', severity='information')
            total_value = annual(series)
            pop = sum(population.get(return_row['code'], {}).values())
            if field == 'total_active_members' and pop and total_value is not None and total_value > pop * 0.6:
                add(return_row, field, 'users_population', total_value, pop,
                    f'{total_value:,.0f} active borrowers equals {total_value / pop:.0%} of ONS mid-2024 population '
                    f'({pop:,}). Non-resident members are possible; check active versus registered members.')
            if field == 'hours_public_computers':
                devices = annual(series_by_field['inventory_computers_in_service'])
                if total_value is not None and devices and total_value > devices * 365 * 12:
                    add(return_row, field, 'computer_capacity', total_value, devices * 365 * 12,
                        f'{total_value:,.0f} hours across {devices:,.0f} year-end PCs/devices is '
                        f'{total_value / devices / 365:.1f} hours/device/day. Above a generous 12-hour daily benchmark; '
                        'check units and fleet changes. Year-end inventory is not an annual average.')

        for total_field, detail_fields in BREAKDOWNS.items():
            total_series = series_by_field[total_field]
            for index in ['Annual', *range(12)]:
                total = annual(total_series) if index == 'Annual' else total_series['monthly'][index]
                details = [annual(series_by_field[field]) if index == 'Annual' else
                           series_by_field[field]['monthly'][index] for field in detail_fields]
                if total is None or not any(value is not None and value > 0 for value in details):
                    continue
                detail_sum = sum(value or 0 for value in details)
                if abs(detail_sum - total) <= max(5, total * 0.01):
                    continue
                if index == 'Annual' and not all(comparable(series_by_field[field]) for field in [total_field, *detail_fields]
                                                if annual(series_by_field[field]) not in (None, 0)):
                    continue
                label = 'Annual' if index == 'Annual' else rotation.MONTHS[index].title()
                overage = detail_sum > total
                if index != 'Annual':
                    active = [total_series, *(series_by_field[field] for field, value in zip(detail_fields, details) if value)]
                    annual_only = any(item['annual'] and sum(value is not None and value > 0 for value in item['monthly']) <= 1
                                      for item in active)
                    if annual_only:
                        continue
                if overage or all(value is not None for value in details):
                    add(return_row, total_field, 'breakdown_sum', total, detail_sum,
                        f'{label}: reported total {total:,.0f}; category sum {detail_sum:,.0f}; '
                        f'difference {total - detail_sum:+,.0f}. '
                        + ('Categories exceed total; check duplication, scope and reporting frequency. Event audiences may overlap.'
                           if overage else 'Remainder may be uncategorised activity; not necessarily a reporting error.'),
                        severity='suspicious' if total > 0 and (overage or total_field == 'total_active_members') else 'information')
            for first, second in combinations(detail_fields, 2):
                first_values, second_values = series_by_field[first]['monthly'], series_by_field[second]['monthly']
                if first_values == second_values and sum(value is not None and value > 0 for value in first_values) >= 6:
                    add(return_row, first, 'identical_categories', annual(series_by_field[first]), annual(series_by_field[second]),
                        f'{first} and {second} are identical across every supplied month. Possible copied series; '
                        'may be deliberate combined-age reporting, so no correction inferred.')

        for field, (_, total_field) in AUXILIARY.items():
            if field not in series_by_field:
                continue
            supplied = annual(series_by_field[field])
            calculated = sum(annual(series_by_field[detail]) or 0 for detail in BREAKDOWNS[total_field])
            if supplied is not None and abs(supplied - calculated) > 1:
                add(return_row, field, 'workbook_summary_sum', supplied, calculated,
                    f'Workbook category-summary cell {supplied:,.0f}; sum of category annual values {calculated:,.0f}. '
                    'Check the workbook formula or independently entered summary.', severity='information')
        impairment = series_by_field.get('active_members_vision_print_impairment')
        if impairment and annual(impairment) and annual(series_by_field['total_active_members']):
            if annual(impairment) > annual(series_by_field['total_active_members']):
                add(return_row, 'active_members_vision_print_impairment', 'subset_exceeds_total',
                    annual(impairment), annual(series_by_field['total_active_members']),
                    'Active users with vision/print impairment exceed all active members. This question is not currently exported to the dashboard.')

        for event_field, attendance_field in [('total_physical_events', 'total_attendees_physical_events'),
                                               ('total_digital_events', 'total_attendees_digital_events')]:
            events = annual(series_by_field[event_field])
            attendance = annual(series_by_field[attendance_field])
            if events and attendance and comparable(series_by_field[event_field]) and comparable(series_by_field[attendance_field]):
                turnout = attendance / events
                if turnout < 1 or turnout > 500:
                    add(return_row, event_field, 'attendance_per_event', events, attendance,
                        f'{events:,.0f} events and {attendance:,.0f} attendees imply {turnout:.2f} attendees/event. '
                        'Check attendance entered as events, coverage differences or large online audiences.')

        for first, second in [
                ('total_physical_audiobook_issues', 'total_digital_audiobook_issues'),
                ('total_ebook_issues', 'total_digital_audiobook_issues'),
                ('total_physical_events', 'total_attendees_physical_events'),
                ('total_digital_events', 'total_attendees_digital_events')]:
            first_series, second_series = series_by_field[first], series_by_field[second]
            values = first_series['monthly']
            if values == second_series['monthly'] and sum(value is not None and value > 0 for value in values) >= 6:
                add(return_row, first, 'identical_metrics', annual(first_series), annual(second_series),
                    f'{first} and {second} are identical month-for-month across every supplied month '
                    f'(annual sums {annual(first_series):,.0f} each). Check copied columns, format classification '
                    'or events entered as attendance; no replacement inferred.')

    history = {}
    for return_row in returns:
        for field, series in return_row['series'].items():
            history.setdefault((return_row['code'], field), []).append(return_row)
    for (_, field), rows in history.items():
        rows.sort(key=lambda row: row['year'])
        for previous, current in zip(rows, rows[1:]):
            if int(current['year'][:4]) - int(previous['year'][:4]) != 1:
                continue
            old, new = previous['series'][field], current['series'][field]
            before, after = annual(old), annual(new)
            if before and after and max(before, after) >= 1000 and comparable(old) and comparable(new):
                ratio = after / before
                if ratio >= 3 or ratio <= 1 / 3:
                    add(current, field, 'year_change', after, before,
                        f'{previous["year"]}: {before:,.0f}; {current["year"]}: {after:,.0f} '
                        f'({ratio:.2f} times; {(ratio - 1) * 100:+.1f}%). Check units, collection scope and completeness; '
                        'a change alone does not prove error.')

    rotation.write_csv(OUTPUT, FIELDS, findings)
    rotation.write_csv(MAPPINGS, list(mappings[0]), mappings)
    rotation.write_csv(RETURNS, list(coverage[0]), coverage)
    rotation.write_csv(EXPLANATIONS, ['Period', 'Authority code', 'Authority name',
                                   'Source column', 'Source row', 'Explanation'], explanations)
    print('Returns:', dict(Counter(row['Period'] for row in coverage)))
    print('Mapped cells checked:', sum(row['Mapped cells'] for row in coverage))
    print('Numeric cells checked:', sum(row['Numeric cells'] for row in coverage))
    print('Findings:', len(findings), dict(Counter(row['Check'] for row in findings)))
    print('Suspicious / information:', dict(Counter(row['Severity'] for row in findings)))
    return findings


def update_register(findings):
    with open(rotation.ERRORS_CSV, newline='', encoding='utf-8') as source:
        reader = csv.DictReader(source)
        fieldnames = reader.fieldnames
        register = list(reader)
    for row in register:
        row['Notes'] = row['Notes'].split(' Formula check: ', 1)[0]
    for finding in findings:
        if finding['Check'] != 'annual_month_sum':
            continue
        evidence = ' Formula check: ' + finding['Field'] + ': ' + finding['Notes']
        related_fields = (finding['Field'], *BREAKDOWNS.get(finding['Field'], ()))
        for row in register:
            if (row['Period'], row['Dataset'], row['Authority code']) != (
                    finding['Period'], finding['Dataset'], finding['Authority code']):
                continue
            if row['Scope'] == 'total' or (row['Scope'] == 'series' and row['Match'] in related_fields):
                if evidence not in row['Notes']:
                    row['Notes'] += evidence
    grouped = {}
    for finding in findings:
        if finding['Severity'] != 'suspicious' or finding['Known issue'] or finding['Field'] in AUXILIARY:
            continue
        check = finding['Check']
        reported, comparator = numeric(finding['Reported']), numeric(finding['Comparator'])
        material = check in ('annual_month_sum', 'identical_categories', 'identical_metrics', 'invalid_count', 'subset_exceeds_total')
        if check == 'breakdown_sum':
            material = (finding['Notes'].startswith('Annual') and reported is not None and comparator is not None
                        and abs(reported - comparator) >= max(100, abs(reported) * 0.02))
        if check == 'monthly_outlier':
            material = comparator is not None and comparator >= 1000 and reported is not None and (
                reported >= comparator * 8 or reported <= comparator * 0.12)
        if check == 'year_change':
            material = (finding['Field'].startswith('total_') or finding['Field'] in
                        ('physical_visits', 'physical_visits_no_colocation', 'hours_public_computers', 'wifi_sessions'))
            material = material and reported is not None and comparator and (
                reported >= comparator * 5 or reported <= comparator * 0.2)
        if check == 'users_population':
            material = finding['Authority name'] != 'City of London'
        if not material or finding['Field'] == 'active_members_vision_print_impairment':
            continue
        key = (finding['Period'], finding['Dataset'], finding['Authority code'])
        grouped.setdefault(key, []).append(finding)

    added = 0
    for (year, dataset, code), issues in grouped.items():
        notes = 'Source audit: ' + ' | '.join(dict.fromkeys(
            f"{issue['Field']}: {issue['Notes']}" for issue in issues))
        notes += ' Values retained as reported pending source confirmation.'
        existing = next((row for row in register if
                         (row['Period'], row['Dataset'], row['Authority code'], row['Scope']) ==
                         (year, dataset, code, 'total')), None)
        if existing:
            if notes not in existing['Notes']:
                existing['Notes'] += ' ' + notes
        elif not any(row['Status'] in ('replaced', 'excluded') and
                     (row['Period'], row['Dataset'], row['Authority code']) == (year, dataset, code)
                     for row in register):
            register.append({'Period': year, 'Dataset': dataset, 'Authority code': code,
                             'Authority name': issues[0]['Authority name'], 'Scope': 'total',
                             'Match': '', 'Status': 'suspicious', 'Estimated count': '', 'Notes': notes})
            added += 1
        affected_fields = dict.fromkeys(
            field for issue in issues
            for field in (issue['Field'], *BREAKDOWNS.get(issue['Field'], ())))
        for field in affected_fields:
            matching = next((row for row in register if
                             (row['Period'], row['Dataset'], row['Authority code'], row['Scope'], row['Match']) ==
                             (year, dataset, code, 'series', field)), None)
            if not matching:
                register.append({'Period': year, 'Dataset': dataset, 'Authority code': code,
                                 'Authority name': issues[0]['Authority name'], 'Scope': 'series',
                                 'Match': field, 'Status': 'suspicious', 'Estimated count': '', 'Notes': notes})
                added += 1

    with open(EXPLANATIONS, newline='', encoding='utf-8') as source:
        explanation_rows = list(csv.DictReader(source))
    contextual = 0
    for row in register:
        for explanation in explanation_rows:
            if (row['Period'], row['Authority code']) != (explanation['Period'], explanation['Authority code']):
                continue
            outreach = explanation['Source column'] in ('Q22', 'Q13') and row['Dataset'] in ('visits', 'click_and_collect')
            if not outreach:
                continue
            text = f" Respondent explanation ({explanation['Source column']}, outreach only): {explanation['Explanation']}"
            if explanation['Explanation'] not in row['Notes']:
                row['Notes'] += text
                contextual += 1
    rotation.write_csv(rotation.ERRORS_CSV, fieldnames, register)
    print(f'Register: {added} new rules; {contextual} relevant respondent explanations attached.')


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--update-register', action='store_true',
                        help='Add material suspicious findings and relevant respondent explanations; never infer corrections.')
    arguments = parser.parse_args()
    audit_findings = audit()
    if arguments.update_register:
        update_register(audit_findings)